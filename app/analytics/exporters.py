import csv
import io
import threading
from datetime import datetime
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import letter, legal, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from app.trainings.services import sanitize_formula_injection
from app.utils.email_utils import get_logo_html, send_email


# ─────────────────────────────────────────────────────────────
# 1. EXPORTADOR A CSV (UTF-8 con BOM y Delimitador ';')
# ─────────────────────────────────────────────────────────────
def export_to_csv(headers, data_rows, filename="reporte_sages.csv"):
    """
    Genera un archivo plano delimitado por ';' con codificación UTF-8 con BOM (utf-8-sig)
    para compatibilidad inmediata con hojas de cálculo (Microsoft Excel en español).
    Sanitiza cada celda contra inyección de fórmulas DDE.
    """
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow(headers)

    for row in data_rows:
        writer.writerow([sanitize_formula_injection(str(cell) if cell is not None else '') for cell in row])

    mem_file = io.BytesIO()
    # utf-8-sig antepone el BOM para que Excel en Windows abra tildes sin corromperse
    mem_file.write(output.getvalue().encode('utf-8-sig'))
    mem_file.seek(0)
    return mem_file, 'text/csv', filename


# ─────────────────────────────────────────────────────────────
# 2. EXPORTADOR A EXCEL (OpenPyXL con Estilos Institucionales)
# ─────────────────────────────────────────────────────────────
def export_to_excel(headers, data_rows, sheet_title="Reporte", filename="reporte_sages.xlsx"):
    """
    Genera un libro de cálculo Excel (.xlsx) con los colores corporativos de CORPOELEC (#0F3669),
    bordes sutiles, auto-ajuste de ancho de columnas y sanitización de fórmulas en cada celda.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(sheet_title)[:31]

    # Paleta corporativa
    header_fill = PatternFill(start_color="0F3669", end_color="0F3669", fill_type="solid")
    header_font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Escribir encabezados
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin

    # Escribir filas sanitizadas
    for row_idx, row in enumerate(data_rows, start=2):
        clean_row = [sanitize_formula_injection(str(c)) if isinstance(c, str) else c for c in row]
        ws.append(clean_row)
        for col_idx in range(1, len(row) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center")

    # Ajuste automático del ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    mem_file = io.BytesIO()
    wb.save(mem_file)
    mem_file.seek(0)
    return mem_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', filename


# ─────────────────────────────────────────────────────────────
# 3. EXPORTADOR A PDF (ReportLab con Matrices y Hoja Adaptable)
# ─────────────────────────────────────────────────────────────



def get_adaptive_page_setup(col_count, force_page_size=None):
    """
    Selecciona automáticamente entre Tamaño Carta (letter: 792x612 pt) y
    Tamaño Oficio (legal: 1008x612 pt) según la densidad de columnas del reporte.
    """
    if force_page_size == 'legal' or (force_page_size is None and col_count > 8):
        page_size = landscape(legal)   # 1008 x 612 pt (Oficio Horizontal)
        page_label = "Formato Oficio Horizontal (Legal — 35.6 x 21.6 cm)"
    else:
        page_size = landscape(letter)  # 792 x 612 pt (Carta Horizontal)
        page_label = "Formato Carta Horizontal (Letter — 27.9 x 21.6 cm)"

    page_w, _ = page_size
    margin_x = 26.0
    usable_width = page_w - (margin_x * 2.0)
    return page_size, page_w, margin_x, usable_width, page_label


def export_to_pdf(headers, data_rows, title="Reporte Institucional", filename="reporte_sages.pdf", user_name="Administrador", report_meta=None):
    """
    Genera un documento PDF formal con Estructura Ejecutiva de Doble Nivel:
    - Banda superior de 4 KPIs macro
    - Sección 1: Matriz Consolidada / Cuadro Resumen (cuando include_metrics está activo)
    - Sección 2: Relación Detallada Estilizada con celdas de doble nivel tipográfico
    - Adaptación automática Carta / Oficio según densidad de columnas
    """
    report_meta = report_meta or {}
    pdf_headers = report_meta.get('pdf_headers') or headers
    pdf_rows = report_meta.get('pdf_rows') or data_rows
    force_page_size = report_meta.get('force_page_size')

    page_size, page_w, margin_x, usable_width, page_label = get_adaptive_page_setup(len(pdf_headers), force_page_size)

    mem_file = io.BytesIO()
    doc = SimpleDocTemplate(
        mem_file,
        pagesize=page_size,
        leftMargin=margin_x,
        rightMargin=margin_x,
        topMargin=22,
        bottomMargin=32
    )
    story = []
    styles = getSampleStyleSheet()
    emitted_at = datetime.now().strftime('%d/%m/%Y %H:%M')
    subtitle_scope = report_meta.get('subtitle_scope', 'Institucional')

    # ── 1. ESTILOS TIPOGRÁFICOS ──
    title_style = ParagraphStyle(
        'HeaderTitle', parent=styles['Heading1'],
        fontName='Helvetica-Bold', fontSize=13,
        textColor=colors.HexColor('#0F3669'), alignment=1, spaceAfter=2
    )
    subtitle_style = ParagraphStyle(
        'HeaderSubtitle', parent=styles['Normal'],
        fontName='Helvetica', fontSize=8.2,
        textColor=colors.HexColor('#475569'), alignment=1, spaceAfter=8
    )
    section_title_style = ParagraphStyle(
        'SectionHeading', parent=styles['Heading2'],
        fontName='Helvetica-Bold', fontSize=9,
        textColor=colors.HexColor('#0F3669'), alignment=0, spaceBefore=2, spaceAfter=4
    )
    kpi_label_style = ParagraphStyle(
        'KpiLabel', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=7.2,
        textColor=colors.HexColor('#64748B'), alignment=1, leading=9
    )
    kpi_value_style = ParagraphStyle(
        'KpiValue', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=10.5,
        textColor=colors.HexColor('#0F3669'), alignment=1, leading=12.5
    )
    th_style = ParagraphStyle(
        'TableHeaderCell', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=7.6,
        textColor=colors.white, alignment=1, leading=9.2
    )
    td_style = ParagraphStyle(
        'TableCell', parent=styles['Normal'],
        fontName='Helvetica', fontSize=7.4,
        textColor=colors.HexColor('#1E293B'), leading=9.2
    )
    td_center_style = ParagraphStyle('TableCellCenter', parent=td_style, alignment=1)
    td_bold_style = ParagraphStyle('TableCellBold', parent=td_style, fontName='Helvetica-Bold')
    td_bold_center_style = ParagraphStyle('TableCellBoldCenter', parent=td_center_style, fontName='Helvetica-Bold')

    # ── 2. ENCABEZADO INSTITUCIONAL ──
    story.append(Paragraph("CORPORACIÓN ELÉCTRICA NACIONAL (CORPOELEC) — DIRECCIÓN UREE", title_style))
    story.append(Paragraph(
        f"<b>SAGES</b> — {escape(title.upper())} &nbsp;|&nbsp; <b>Alcance:</b> {escape(subtitle_scope)} &nbsp;|&nbsp; "
        f"<b>Emitido por:</b> {escape(user_name)} &nbsp;|&nbsp; <b>Fecha:</b> {emitted_at}",
        subtitle_style
    ))

    # ── 3. BANDA EJECUTIVA DE KPIS Y SECCIÓN 1 (MATRIZ CONSOLIDADA) ──
    if report_meta.get('include_metrics', True) and report_meta.get('kpi_headers'):
        kpi_headers = [Paragraph(escape(h), kpi_label_style) for h in report_meta['kpi_headers']]
        kpi_values = [Paragraph(escape(v), kpi_value_style) for v in report_meta['kpi_values']]
        kpi_table = Table([kpi_headers, kpi_values], colWidths=[usable_width / 4.0] * 4)
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0, 0), (-1, -1), 4.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5.5),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE')
        ]))
        story.append(kpi_table)
        story.append(Spacer(1, 7))

        # Renderizar Sección 1: Matriz Consolidada
        sum_headers = report_meta.get('summary_headers', [])
        sum_rows = report_meta.get('summary_rows', [])
        sum_weights = report_meta.get('summary_weights', [1.0] * len(sum_headers))
        center_sum = report_meta.get('center_cols_summary', set())

        if sum_headers and sum_rows:
            story.append(Paragraph(escape(report_meta.get('summary_title', '1. Resumen Consolidado')), section_title_style))
            total_w_sum = sum(sum_weights) or 1.0
            sum_widths = [(w / total_w_sum) * usable_width for w in sum_weights]

            sum_table_data = [[Paragraph(escape(h), th_style) for h in sum_headers]]
            for r_idx, row in enumerate(sum_rows):
                is_total = (r_idx == len(sum_rows) - 1 and 'total' in str(row[0]).lower())
                formatted_row = []
                for c_idx, val in enumerate(row):
                    c_name = sum_headers[c_idx]
                    st = (td_bold_center_style if c_name in center_sum else td_bold_style) if is_total else (td_center_style if c_name in center_sum else td_style)
                    formatted_row.append(Paragraph(escape(str(val)), st))
                sum_table_data.append(formatted_row)

            sum_table = Table(sum_table_data, colWidths=sum_widths, repeatRows=1)
            sum_styles = [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 3.8),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3.8),
                ('LEFTPADDING', (0, 0), (-1, -1), 5),
                ('RIGHTPADDING', (0, 0), (-1, -1), 5),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
            ]
            if len(sum_rows) > 1 and 'total' in str(sum_rows[-1][0]).lower():
                sum_styles.append(('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E2E8F0')))
            sum_table.setStyle(TableStyle(sum_styles))
            story.append(sum_table)
            story.append(Spacer(1, 8))

    # ── 4. SECCIÓN 2: TABLA DE DETALLE ESTILIZADA ──
    if report_meta.get('detail_title'):
        story.append(Paragraph(escape(report_meta['detail_title']), section_title_style))

    pdf_weights = report_meta.get('pdf_weights') or [1.0] * len(pdf_headers)
    total_w_det = sum(pdf_weights) or 1.0
    det_widths = [(w / total_w_det) * usable_width for w in pdf_weights]
    center_det = report_meta.get('center_cols_detail', set())

    det_table_data = [[Paragraph(escape(h), th_style) for h in pdf_headers]]
    for row in pdf_rows:
        formatted_row = []
        for c_idx, cell_val in enumerate(row):
            col_name = pdf_headers[c_idx] if c_idx < len(pdf_headers) else ''
            st = td_center_style if col_name in center_det else td_style
            parts = [escape(p) for p in str(cell_val)[:200].split('\n')]
            if len(parts) > 1:
                html_cell = f"<b>{parts[0]}</b><br/><font color='#475569' size='6.8'>{parts[1]}</font>"
            else:
                html_cell = parts[0]
            formatted_row.append(Paragraph(html_cell, st))
        det_table_data.append(formatted_row)

    det_table = Table(det_table_data, colWidths=det_widths, repeatRows=1)
    det_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F3669')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3.8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.8),
        ('LEFTPADDING', (0, 0), (-1, -1), 4.5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4.5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))
    story.append(det_table)

    # ── 5. PIE DE PÁGINA OFICIAL Y NUMERACIÓN ──
    def _draw_footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor('#CBD5E1'))
        canvas.setLineWidth(0.5)
        canvas.line(margin_x, 24, page_w - margin_x, 24)
        canvas.setFont('Helvetica', 7.5)
        canvas.setFillColor(colors.HexColor('#64748B'))
        canvas.drawString(margin_x, 13, f"SAGES — CORPOELEC (UREE) | {page_label} | Generado: {emitted_at}")
        canvas.drawRightString(page_w - margin_x, 13, f"Página {doc_obj.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_draw_footer, onLaterPages=_draw_footer)
    mem_file.seek(0)
    return mem_file, 'application/pdf', filename


# ─────────────────────────────────────────────────────────────
# 4. DESPACHO ASÍNCRONO POR CORREO ELECTRÓNICO (ATTACHMENT)
# ─────────────────────────────────────────────────────────────
def send_report_email_async(app, to_email, user_name, report_title, file_bytes, filename, mimetype):
    """
    Despacha el archivo generado adjunto al correo del usuario en un hilo desacoplado.
    Usa send_email() centralizado con el template HTML corporativo estándar de SAGES.
    """
    import html as html_lib
    safe_name = html_lib.escape(str(user_name))
    safe_title = html_lib.escape(str(report_title))
    safe_filename = html_lib.escape(str(filename))

    # Determinar etiqueta del formato
    ext = filename.rsplit('.', 1)[-1].upper() if '.' in filename else 'ARCHIVO'
    emitted_at = datetime.now().strftime('%d/%m/%Y a las %H:%M')

    subject = f"Reporte Estadístico: {report_title} — SAGES"

    text_content = (
        f"Estimado/a {user_name},\n\n"
        f"Se ha generado exitosamente el reporte institucional solicitado: '{report_title}'.\n\n"
        f"Adjunto a este correo encontrará el documento en formato {ext}.\n\n"
        f"Generado el {emitted_at} por el Sistema SAGES — CORPOELEC UREE.\n\n"
        f"Por favor, no responda a este correo automatizado."
    )

    logo_html = get_logo_html()
    html_content = f"""<!DOCTYPE html>
    <html lang="es">
    <head>
        <meta charset="UTF-8">
        <title>{subject}</title>
    </head>
    <body style="margin: 0; padding: 0; font-family: Arial, sans-serif; background-color: #f9fafb; color: #1f2937;">
        <table align="center" border="0" cellpadding="0" cellspacing="0" width="100%" style="max-width: 600px; margin: 20px auto; background-color: #ffffff; border-radius: 16px; box-shadow: 0 4px 12px rgba(0, 0, 0, 0.05); border: 1px solid #e5e7eb; overflow: hidden;">
            <tr>
                <td style="height: 6px; background: linear-gradient(90deg, #1c3d73 0%, #1fcab0 100%);"></td>
            </tr>
            <tr>
                <td align="center" style="padding: 40px 20px 20px 20px;">
                    {logo_html}
                    <h1 style="font-size: 15px; color: #4b5563; margin: 0 0 25px 0; line-height: 1.5;">Sistema de Gestión de Solicitudes de Formación</h1>
                </td>
            </tr>
            <tr>
                <td align="center" style="padding: 0 40px;">
                    <h2 style="font-size: 22px; font-weight: 700; color: #1f2937; margin: 10px 0 20px 0; text-transform: uppercase; letter-spacing: 0.5px;">REPORTE ESTADÍSTICO GENERADO</h2>
                    <p style="font-size: 15px; color: #4b5563; margin: 0 0 10px 0; line-height: 1.5; text-align: left;">
                        Estimado/a <strong>{safe_name}</strong>,
                    </p>
                    <p style="font-size: 15px; color: #4b5563; margin: 0 0 20px 0; line-height: 1.6; text-align: left;">
                        Le informamos que el reporte estadístico solicitado ha sido <strong>generado exitosamente</strong> y se encuentra adjunto al presente correo.
                    </p>
                </td>
            </tr>
            <tr>
                <td style="padding: 0 40px 20px 40px;">
                    <div style="background-color: #f3f4f6; border-radius: 8px; padding: 16px 20px; border: 1px solid #e5e7eb;">
                        <table width="100%" cellpadding="4" cellspacing="0" style="font-size: 14px;">
                            <tr>
                                <td style="color: #6b7280; font-weight: 600; width: 35%;">Reporte:</td>
                                <td style="color: #1c3d73; font-weight: 700;">{safe_title}</td>
                            </tr>
                            <tr>
                                <td style="color: #6b7280; font-weight: 600;">Archivo:</td>
                                <td style="color: #111827;">{safe_filename}</td>
                            </tr>
                            <tr>
                                <td style="color: #6b7280; font-weight: 600;">Formato:</td>
                                <td style="color: #111827;">{ext}</td>
                            </tr>
                            <tr>
                                <td style="color: #6b7280; font-weight: 600;">Generado:</td>
                                <td style="color: #111827;">{emitted_at}</td>
                            </tr>
                        </table>
                    </div>
                </td>
            </tr>
            <tr>
                <td align="center" style="padding: 10px 40px 40px 40px;">
                    <hr style="border: 0; border-top: 1px solid #e5e7eb; margin-bottom: 25px;">
                    <p style="font-size: 12px; color: #9ca3af; line-height: 1.6; margin: 0; text-align: justify;">
                        * Este mensaje fue generado de forma automática por el módulo de estadísticas de SAGES. No responda a este correo.
                    </p>
                </td>
            </tr>
            <tr>
                <td align="center" style="background-color: #f9fafb; padding: 20px; font-size: 11px; color: #9ca3af; border-top: 1px solid #e5e7eb;">
                    &copy; 2026 SAGES — CORPOELEC. Todos los derechos reservados.
                </td>
            </tr>
        </table>
    </body>
    </html>"""

    # Preparar adjunto y despachar con send_email() centralizado
    maintype, subtype = mimetype.split('/', 1)
    attachments = [{
        'filename': filename,
        'maintype': maintype,
        'subtype': subtype,
        'content': file_bytes.getvalue()
    }]

    def _worker():
        with app.app_context():
            ok = send_email(to_email, subject, html_content, text_content, attachments)
            if ok:
                print(f"[SMTP ÉXITO]: Reporte '{filename}' enviado a {to_email}")
            else:
                print(f"[ERROR SMTP]: Falló el envío del reporte '{filename}' a {to_email}")

    threading.Thread(target=_worker, daemon=True).start()
